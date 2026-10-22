"""Builds docs/testing/ATS_TestCases_Master.xlsx.

Carries forward the 118 cases from 'ATS_TestCases_HRUpload_to_ZekoHR_Updated 2.xlsx'
(with corrections where the code has since changed) and adds new modules.
"""
import sys
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

SRC = sys.argv[1]
OUT = sys.argv[2]

# Automation classes
A_UI = "Automated - UI"
A_UIF = "Automated - UI (fixture)"
A_API = "Automated - API"
M_EXT = "Manual - External"
M_JOB = "Manual - Scheduled job"
M_EXP = "Manual - Exploratory"

EXT_AI_PARSE = "AI resume parser (Gemini / OpenRouter / Cohere)"
EXT_AI_SEARCH = "AI embeddings (Gemini) for JD / keyword search"
EXT_ZEKO = "Zeko API"
EXT_MAIL = "MS Graph mail (delivery check only)"
EXT_TEAMS = "MS Graph calendar / Teams"
EXT_ONEDRIVE = "MS Graph OneDrive"
EXT_TURNSTILE = "Cloudflare Turnstile"

HEADERS = [
    "TC ID", "Module", "Feature", "Test Scenario", "Preconditions / Test Data",
    "Test Steps", "Expected Result", "Type", "Priority", "Role",
    "Automation", "External Dependency", "Source", "Previous Run (2026-08-11)",
    "Status", "Actual Result", "Remarks",
]

# ─────────────────────────────────────────────────────────────────────────────
# 1. Existing cases — automation class per ID prefix, then per-ID overrides.
# ─────────────────────────────────────────────────────────────────────────────
PREFIX_DEFAULTS = {
    "HRU": dict(automation=M_EXT, ext=EXT_AI_PARSE, role="Recruiter", priority="P1"),
    "SCR": dict(automation=M_EXT, ext=EXT_AI_SEARCH, role="Recruiter", priority="P1"),
    "CPIPE": dict(automation=A_UIF, ext="", role="Recruiter", priority="P1"),
    "ZHR": dict(automation=M_EXT, ext=EXT_ZEKO, role="Recruiter", priority="P2"),
    "OUT": dict(automation=A_UIF, ext=EXT_MAIL, role="Recruiter", priority="P1"),
    "POST": dict(automation=A_UIF, ext="", role="Recruiter", priority="P2"),
    "ASM": dict(automation=A_UIF, ext=EXT_MAIL, role="Recruiter", priority="P1"),
    "ZFN": dict(automation=A_UIF, ext=EXT_MAIL, role="Recruiter", priority="P1"),
    "T1": dict(automation=A_UIF, ext=EXT_MAIL, role="Recruiter", priority="P1"),
    "T2": dict(automation=A_UIF, ext=EXT_MAIL, role="Recruiter", priority="P2"),
    "T3": dict(automation=A_UIF, ext=EXT_MAIL, role="Recruiter", priority="P2"),
    "HRR": dict(automation=A_UIF, ext=EXT_MAIL, role="Recruiter", priority="P2"),
    "CEO": dict(automation=A_UIF, ext=EXT_MAIL, role="Recruiter", priority="P2"),
    "CLI": dict(automation=A_UIF, ext=EXT_MAIL, role="Recruiter", priority="P2"),
    "DOC": dict(automation=A_UIF, ext=EXT_MAIL, role="Recruiter", priority="P1"),
    "OFR": dict(automation=A_UIF, ext="", role="Recruiter", priority="P1"),
}

MODULE_GROUP = {
    "HRU": "HR Manual Upload", "SCR": "Candidate Screening", "CPIPE": "Candidate Pipeline",
    "ZHR": "Zeko HR Screening", "OUT": "Candidate Pipeline", "POST": "Candidate Pipeline",
    "ASM": "Candidate Pipeline", "ZFN": "Candidate Pipeline", "T1": "Candidate Pipeline",
    "T2": "Candidate Pipeline", "T3": "Candidate Pipeline", "HRR": "Candidate Pipeline",
    "CEO": "Candidate Pipeline", "CLI": "Candidate Pipeline", "DOC": "Documents",
    "OFR": "Offer & Closure",
}

# Per-ID changes. Keys: scenario, steps, expected, type, automation, ext, priority,
# remarks_add (appended to Remarks), pre.
OVERRIDES = {
    "HRU-01": dict(remarks_add="Total Experience defect fixed in CHANGES-2026-08-07 #19. Re-test with a CV using 'Jun-2022 - Present', then Search Candidate > View Details."),
    "HRU-02": dict(remarks_add="Scenario title says 'exceeding max count/size' but steps/expected describe a multi-file success. Re-word before next run. Total Experience fix: #19."),
    "HRU-06": dict(remarks_add="Never executed (no Actual Result in the 2026-08-11 pass)."),
    "HRU-07": dict(remarks_add="Never executed. Missing-data link also exercised by MJD-01..04."),
    "HRU-08": dict(remarks_add="Never executed."),
    "HRU-09": dict(automation=M_EXT, remarks_add="Never executed. Live socket update itself is automatable once a job can be produced without the AI parser."),
    "SCR-05": dict(remarks_add="Validation itself is client-side; only reaching the result list needs AI search."),
    "CPIPE-01": dict(pre="Journey created by shortlisting via API (fixture) because the screening search needs AI embeddings."),
    "CPIPE-02": dict(pre="Journey created by shortlisting via API (fixture)."),
    "ZHR-04": dict(expected="Cancellation succeeds without a reason - the reason field is optional. Pipeline row status = 'cancelled'.",
                   remarks_add="Expected result clarified (was 'As it's Optional')."),
    "ZHR-05": dict(automation=M_JOB, ext=EXT_ZEKO),
    "ZHR-06": dict(automation=M_JOB, ext=EXT_ZEKO, remarks_add="Never executed."),
    "OUT-01": dict(remarks_add="Locally the 'Emails in this round' log reads 'failed' (no Graph). Inbox delivery is a manual check on staging."),
    "OUT-06": dict(ext=""),
    "ASM-02": dict(automation=A_UIF, ext=EXT_MAIL, remarks_add="Modal + link entry automated; inbox delivery manual."),
    "ASM-03": dict(automation=A_UIF, ext=EXT_MAIL),
    "ASM-04": dict(automation=A_UIF, ext="", scenario="Assessment scores imported from Evalground CSV",
                   steps="1. At IQ / Tech Assessment, click Upload / Import\n2. Choose a valid Evalground CSV containing the candidate's email\n3. Confirm the import",
                   expected="Rows matched to candidates by email; scores written to the candidate's assessment record; unmatched rows listed; optional auto approve/reject on pass mark if enabled in Settings",
                   remarks_add="Steps/expected expanded from 'Upload csv manually'. Test data: docs/*TestReport*.csv."),
    "ASM-09": dict(ext=""),
    "ZFN-02": dict(automation=M_EXT, ext=EXT_ZEKO, scenario="Assign & schedule Zeko functional interview",
                   steps="Same flow as ZHR-01..04 but with the Functional Screening Zeko job",
                   expected="Same behaviour as ZHR-01..04 for the functional round",
                   remarks_add="Original row had no steps/expected."),
    "ZFN-07": dict(ext=""),
    "T1-02": dict(ext=""), "T1-03": dict(ext="", type="Negative"), "T1-04": dict(ext=""),
    "T1-05": dict(ext=""), "T1-06": dict(ext=""),
    "T1-07": dict(ext=EXT_MAIL, remarks_add="Preview text must already show 'Hi <interviewer name>,' (fix #20)."),
    "T1-08": dict(automation=M_EXT, ext=EXT_TEAMS),
    "T1-09": dict(automation=M_EXT, ext=EXT_MAIL, remarks_add="Interviewer-name defect fixed in CHANGES-2026-08-07 #20; needs re-run of prisma/seed-email-templates.js."),
    "T1-10": dict(ext="", pre="MS_CALENDAR_ENABLED=false locally: booking is saved without a Teams event."),
    "T1-11": dict(ext="", steps="After Create invite, re-open the drawer and check the 'Awaiting Results' step"),
    "T1-12": dict(type="Positive", ext=EXT_TEAMS,
                  expected="Modal reopens pre-filled with current date/time/interviewer; saving updates the booking and re-sends invites. The SAME Teams/Outlook event is updated (time patched) - no new meeting ID is created.",
                  remarks_add="'Verify' resolved from code: rescheduleInterviewRound() calls updateInterviewEventTime() on the existing graph_event_id."),
    "T1-13": dict(type="Positive", ext=EXT_TEAMS,
                  expected="'Confirm Cancel Interview' modal opens; 'Cancel reason (optional)' may be left blank; booking status = cancelled; Teams event cancelled; cancellation notices sent (mailbox check manual)",
                  remarks_add="'Verify' resolved from code: cancel reason is optional."),
    "T1-14": dict(ext=""),
    "T1-15": dict(automation=M_EXP, ext="", remarks_add="Open product question - confirm with RT whether outcome before interview is intended."),
    "T1-20": dict(ext=""), "T2-07": dict(ext=""), "T3-07": dict(ext=""), "HRR-06": dict(ext=""),
    "CEO-06": dict(ext=""), "CLI-06": dict(ext=""), "DOC-07": dict(ext=""),
    "T2-02": dict(remarks_add="Covered by repeating T1-02..T1-15 on Technical Round 2."),
    "T3-02": dict(remarks_add="Covered by repeating T1-02..T1-15 on Technical Round 3."),
    "T3-03": dict(expected="Advances to 'HR Round'; label = 'Technical Round 3 Approved'; email sent",
                  remarks_add="Expected result was a copy of the steps; corrected."),
    "T3-08": dict(steps="1. Open a candidate at Technical Round 3\n2. Click the optional-round skip control\n3. Confirm 'Skip this optional round?' > 'Yes, skip it'",
                  expected="Candidate moves to 'HR Round'; a distinct 'skip' event is logged, separate from a normal approve event",
                  remarks_add="Skip is now available in the drawer UI (was API-only)."),
    "CLI-07": dict(steps="1. Open a candidate at Client Round\n2. Click the optional-round skip control\n3. Confirm 'Skip this optional round?' > 'Yes, skip it'",
                   expected="Candidate moves to 'Documents'; distinct 'skip' event logged",
                   remarks_add="Skip is now available in the drawer UI (was API-only)."),
    "DOC-02": dict(steps="1. Open a candidate at Documents\n2. Click 'Send document request'",
                   expected="Request created with a no-login upload link; Documents segment shows the request as sent; email goes to the candidate ONLY (never the vendor)",
                   remarks_add="Inbox delivery is manual; request creation automated."),
    # Offer stage is record-only since 2026-08-25 (CHANGES-2026-08-25-client-round-and-offer-mark-only.md).
    "OFR-01": dict(expected="Card moves to 'Offer' column; the Offer panel shows 'Record-only: HR prepares and shares the letter from its own mailbox' and a 'Record offer shared' button. The generic Approve/Reject/Hold bar is NOT shown on this stage.",
                   remarks_add="Never executed. Expected result updated for the record-only Offer stage (2026-08-25)."),
    "OFR-02": dict(scenario="Record offer shared",
                   steps="1. Click 'Record offer shared'\n2. Pick a proposed joining date (past dates are disabled); add optional remarks\n3. Click 'Record it'",
                   expected="Toast 'Offer recorded as shared.'; panel reads 'Shared <date> · proposed joining <date> · awaiting the candidate's decision'; 'Mark accepted' / 'Mark rejected' buttons appear",
                   remarks_add="Never executed. REPLACES 'Approve (last stage)' - Offer no longer has Approve/Reject/Hold."),
    "OFR-03": dict(scenario="Mark offer accepted", type="Positive",
                   steps="1. After OFR-02, click 'Mark accepted'",
                   expected="Panel reads '... candidate accepted'; offer_accepted_at populated; if accepted hires now equal the MRF's openings, the requisition auto-closes (filled) and leaves JD filtering",
                   remarks_add="Never executed. REPLACES 'Reject without a valid reason'."),
    "OFR-04": dict(scenario="Mark offer rejected (declined by candidate)", type="Positive",
                   steps="1. On a fresh journey, record offer shared\n2. Click 'Mark rejected'",
                   expected="Panel reads '... candidate rejected'; Accepted / Declined row shows 'Declined by the candidate'; MRF seat not filled",
                   remarks_add="Never executed. REPLACES 'Reject with a valid reason'."),
    "OFR-05": dict(scenario="Close candidate record as 'Joined' from Offer",
                   steps="1. After OFR-03, click 'Close candidate record…'\n2. Pick Final status = Joined\n3. Click 'Close the record'",
                   expected="Alert states Joined does not email the candidate; card leaves the board; drawer closes; pipeline_status = 'hired'; joined_at = closed_at; vendor 90-day lock frozen",
                   remarks_add="Never executed. REPLACES 'Hold with a valid reason'."),
    "OFR-06": dict(scenario="Internal approval controls are absent (mark-only regression)",
                   steps="1. Open a candidate at Offer before anything is recorded",
                   expected="No 'Request internal approval' or 'Mark approved' buttons; progress strip does not stick on 'Internal approval not requested yet'",
                   type="Negative",
                   remarks_add="Never executed. REPLACES 'Preview outcome email' - Offer has no outcome email."),
}


def load_existing(path):
    wb = openpyxl.load_workbook(path)
    ws = wb["Test Cases"]
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        tcid = (r[0] or "").strip() if isinstance(r[0], str) else r[0]
        if not tcid or "-" not in str(tcid) or str(tcid).startswith("Feedback") or r[1] is None:
            continue
        rows.append(dict(id=tcid, module=r[1], scenario=r[2], steps=r[3], expected=r[4],
                         type=r[5], actual=r[6], remarks=r[7]))
    return rows


def existing_case_rows(path):
    out = []
    for c in load_existing(path):
        prefix = c["id"].split("-")[0]
        d = PREFIX_DEFAULTS[prefix]
        o = OVERRIDES.get(c["id"], {})
        remarks = c["remarks"] or ""
        if o.get("remarks_add"):
            remarks = (remarks + "\n" if remarks else "") + o["remarks_add"]
        out.append([
            c["id"], MODULE_GROUP[prefix], c["module"],
            o.get("scenario", c["scenario"]),
            o.get("pre", ""),
            o.get("steps", c["steps"]),
            o.get("expected", c["expected"]),
            o.get("type", c["type"]),
            o.get("priority", d["priority"]),
            d["role"],
            o.get("automation", d["automation"]),
            o.get("ext", d["ext"]),
            "ATS_TestCases_HRUpload_to_ZekoHR_Updated 2.xlsx",
            c["actual"] or "Not executed",
            "Not Run", "", remarks,
        ])
    return out


# ─────────────────────────────────────────────────────────────────────────────
# 2. New cases.
# (id, module, feature, scenario, pre, steps, expected, type, priority, role, automation, ext, source)
# ─────────────────────────────────────────────────────────────────────────────
N = []


def add(*a):
    N.append(list(a))


# ── Authentication & Access ──────────────────────────────────────────────────
M = "Authentication & Access"
add("AUTH-01", M, "Login", "Login with valid recruiter credentials", "Active recruiter account",
    "1. Open /login\n2. Enter username/email and password\n3. Click Sign in",
    "Redirected to /dashboard; header shows the user; sidebar shows only modules enabled for the user", "Positive", "P1", "Recruiter", A_UI, "", "auth.routes.js")
add("AUTH-02", M, "Login", "Login with wrong password", "Active account",
    "1. Enter a valid username with a wrong password\n2. Submit",
    "Error 'Invalid username/email or password.'; stays on /login; no token stored", "Negative", "P1", "Any", A_UI, "", "auth.service.js")
add("AUTH-03", M, "Login", "Login with unknown username", "",
    "1. Enter a username that does not exist\n2. Submit",
    "Same generic error 'Invalid username/email or password.' (no hint that the user does not exist)", "Security", "P2", "Any", A_UI, "", "auth.service.js")
add("AUTH-04", M, "Login", "Login with empty fields", "",
    "1. Leave username and/or password blank\n2. Submit",
    "Form validation blocks submit; no API call (or API returns 'Please provide username and password.')", "Negative", "P2", "Any", A_UI, "", "auth.controller.js")
add("AUTH-05", M, "Login", "Deactivated account cannot log in", "Account deactivated via ADM-07",
    "1. Log in with the deactivated account",
    "Error 'Your account has been deactivated. Contact an administrator.'", "Negative", "P1", "Any", A_UI, "", "auth.service.js")
add("AUTH-06", M, "Route guard", "Protected page without a session", "Logged out",
    "1. Open /pipeline (or /mrf, /candidates) directly",
    "Redirected to /login; after login the user can open the page", "Security", "P1", "Any", A_UI, "", "App.jsx ProtectedRoute")
add("AUTH-07", M, "Logout", "Logout ends the session", "Logged in",
    "1. User menu > Logout\n2. Press browser Back",
    "Returned to /login; Back does not show protected content; API calls with the old token are refused", "Security", "P1", "Any", A_UI, "", "MainLayout.jsx")
add("AUTH-08", M, "Change password", "Change password with wrong current password", "Logged in",
    "1. User menu > Change Password\n2. Enter a wrong current password and a new password\n3. Submit",
    "Error 'Current password is incorrect.'; password unchanged", "Negative", "P2", "Any", A_UI, "", "auth.controller.js")
add("AUTH-09", M, "Change password", "Change password successfully", "Logged in (use a disposable test user)",
    "1. Change Password with correct current password\n2. Logout\n3. Login with the new password",
    "Success message; login with the new password works; old password rejected", "Positive", "P2", "Any", A_UI, "", "auth.controller.js")
add("AUTH-10", M, "Forgot password", "Request a password reset", "",
    "1. /login > Forgot password\n2. Enter a registered email\n3. Submit",
    "Generic confirmation shown (does not reveal whether the email exists); reset email sent (inbox check manual)", "Positive", "P2", "Any", A_UI, EXT_MAIL, "ForgotPassword.jsx")
add("AUTH-11", M, "Reset password", "Reset link invalid / tampered", "",
    "1. Open /reset-password?token=invalid\n2. Set a new password",
    "Error 'This password reset link is invalid. Please request a new one.'", "Negative", "P2", "Any", A_UI, "", "auth.service.js")
add("AUTH-12", M, "Reset password", "Reset link already used / expired", "Token used once, or expired (fixture)",
    "1. Re-open a used reset link and submit",
    "Error 'This password reset link has already been used...' or '...has expired...'", "Negative", "P3", "Any", A_UIF, "", "auth.service.js")
add("AUTH-13", M, "Admin login", "Admin portal login", "Superadmin account",
    "1. Open /admin/login\n2. Log in as superadmin",
    "Lands on /admin/dashboard (Admin Portal)", "Positive", "P1", "Superadmin", A_UI, "", "AdminLogin.jsx")
add("AUTH-14", M, "Role guard", "Recruiter cannot open the Admin Portal", "Recruiter logged in",
    "1. Open /admin/dashboard directly",
    "Access denied / redirected; no user list is loaded", "Security", "P1", "Recruiter", A_UI, "", "App.jsx AdminRoute")
add("AUTH-15", M, "Vendor confinement", "Vendor is confined to vendor pages", "Vendor logged in",
    "1. Check the sidebar\n2. Type /pipeline, /mrf, /analytics, /candidates in the address bar",
    "Sidebar shows only 'Dashboard' and 'Upload Candidate'; every other URL redirects to /vendor-dashboard", "Security", "P1", "Vendor", A_UI, "", "MainLayout.jsx VENDOR_ALLOWED_PATHS")
add("AUTH-16", M, "Vendor confinement", "Vendor token is refused by staff APIs", "Vendor JWT",
    "1. Call GET /api/pipeline, /api/screening/roles, /api/hr-upload/... with the vendor token",
    "403 on every staff router (requireStaff), even if a module toggle was wrongly enabled for the vendor", "Security", "P1", "Vendor", A_API, "", "CHANGES-phase3-m6-vendor.md")
add("AUTH-17", M, "Module access", "Disabled module is hidden and guarded", "Recruiter with 'Candidate Pipeline' module disabled (ADM-10)",
    "1. Log in as that recruiter\n2. Check sidebar\n3. Open /pipeline directly",
    "'Candidate Pipeline' not in sidebar; /pipeline is blocked by ModuleRoute", "Security", "P1", "Recruiter", A_UI, "", "App.jsx ModuleRoute")
add("AUTH-18", M, "Captcha", "Login captcha when Turnstile is enabled", "VITE_TURNSTILE_SITE_KEY + secret configured",
    "1. Open /login\n2. Submit without completing the captcha",
    "Error 'Captcha verification is required. Please refresh the page and try again.'", "Negative", "P3", "Any", M_EXT, EXT_TURNSTILE, "auth.controller.js")

# ── Admin Portal ─────────────────────────────────────────────────────────────
M = "Admin Portal"
add("ADM-01", M, "Users", "Create a recruiter user", "Superadmin logged in; at least one company",
    "1. Admin Portal > Users > Add user\n2. Enter email, name, role = Recruiter, company, password\n3. Save",
    "User appears in the list as Active with role Recruiter; can log in (AUTH-01)", "Positive", "P1", "Superadmin", A_UI, "", "admin.controller.js")
add("ADM-02", M, "Users", "Duplicate email is flagged", "User from ADM-01 exists",
    "1. Add user with the same email",
    "Email-availability check flags the address as taken; user not created", "Negative", "P2", "Superadmin", A_UI, "", "admin.routes.js check-email")
add("ADM-03", M, "Users", "Company is mandatory for company roles", "",
    "1. Add user with role Admin/Recruiter/Vendor and no company",
    "Error 'A company is required for this role.'", "Negative", "P2", "Superadmin", A_UI, "", "admin.controller.js")
add("ADM-04", M, "Users", "Company admin cannot create a Super Admin", "Logged in as company Admin",
    "1. Add user - open the Role dropdown\n2. (API) POST /api/admin/users/create with role superadmin",
    "'Super Admin' not offered; API returns 'Admins cannot create SuperAdmin accounts.' (403)", "Security", "P1", "Admin", A_UI, "", "admin.controller.js")
add("ADM-05", M, "Users", "Edit a user", "",
    "1. Select a user > Edit\n2. Change name / company\n3. Save",
    "'User updated successfully.'; list reflects the change", "Positive", "P2", "Superadmin", A_UI, "", "AdminDashboard.jsx")
add("ADM-06", M, "Users", "Cannot change own role", "",
    "1. Edit your own user and change the role",
    "Error 'You cannot change your own role.'", "Negative", "P2", "Admin", A_UI, "", "admin.controller.js")
add("ADM-07", M, "Users", "Deactivate and reactivate a user", "",
    "1. Toggle a recruiter to Inactive\n2. Try to log in as them (AUTH-05)\n3. Toggle back to Active",
    "Status shows Inactive and login is refused; after reactivation login works", "Positive", "P1", "Superadmin", A_UI, "", "admin.controller.js")
add("ADM-08", M, "Users", "Cannot change own status", "",
    "1. Toggle your own account's status",
    "Error 'You cannot change the status of your own account.'", "Negative", "P3", "Admin", A_UI, "", "admin.controller.js")
add("ADM-09", M, "Users", "Delete a user (superadmin only)", "Disposable user",
    "1. As Superadmin, delete the user\n2. As company Admin, try to delete a user",
    "Superadmin: 'User deleted permanently.'; Admin: 'Only a SuperAdmin can delete users.'", "Security", "P2", "Superadmin / Admin", A_UI, "", "admin.controller.js")
add("ADM-10", M, "Module access", "Toggle a module for a user", "Recruiter user",
    "1. Select the user > Module Access\n2. Switch off 'Candidate Pipeline'\n3. Log in as the user",
    "Toggle saved; module hidden and guarded for that user (AUTH-17); switching it back on restores access", "Positive", "P1", "Superadmin", A_UI, "", "admin.routes.js modules/set-access")
add("ADM-11", M, "Users", "Filter the user list", "Users of several roles/companies",
    "1. Filter by Role, Status (Active/Inactive), Company",
    "List narrows to matching users; 'All' restores the full list", "Positive", "P3", "Superadmin", A_UI, "", "AdminDashboard.jsx")
add("ADM-12", M, "Users", "Export users CSV", "",
    "1. Click Export on the user list",
    "CSV downloads with one row per user", "Positive", "P3", "Superadmin", A_UI, "", "admin.routes.js users/export")
add("ADM-13", M, "Companies", "Create / edit / deactivate a company", "Superadmin",
    "1. Companies > Add company (name, slug, domain)\n2. Edit it\n3. Toggle status",
    "'Company created.' / 'Company updated.'; status change reflected", "Positive", "P3", "Superadmin", A_UI, "", "AdminDashboard.jsx")
add("ADM-14", M, "Referral audit", "Referral audit log and export", "At least one referral set/removed (CAND-11/12)",
    "1. Open the referral log\n2. Export",
    "Each set/remove event listed with who, when and reason; CSV downloads", "Positive", "P3", "Admin", A_UI, "", "admin.routes.js referral-log")

# ── Dashboard & Notifications ────────────────────────────────────────────────
M = "Dashboard & Notifications"
add("DASH-01", M, "Dashboard", "Dashboard loads stats and recent uploads", "Seeded candidates",
    "1. Log in as recruiter\n2. Observe /dashboard",
    "KPI cards, recent uploads and recruiter breakdown load without errors", "Positive", "P2", "Recruiter", A_UI, "", "dashboard.routes.js")
add("DASH-02", M, "Dashboard", "Quick Actions navigate", "",
    "1. Click each Quick Action (Candidate Screening, Recruitment Analytics, New MRF Request, Search & Edit Candidates, HR Manual Upload, Vendor Upload, System Configuration)",
    "Each opens the matching page", "Positive", "P3", "Recruiter", A_UI, "", "Dashboard.jsx")
add("DASH-03", M, "Dashboard", "Shortlist KPI excludes closed candidates", "A candidate closed as Joined (OFR-05)",
    "1. Compare the Shortlisted KPI before and after closing a candidate",
    "Closed candidate no longer counts as Shortlisted; recruiter leaderboard credit is kept (not deducted on hire)", "Positive", "P3", "Recruiter", A_UIF, "", "CHANGES-2026-08-26-candidate-closure-graceful-exit.md")
add("NTF-01", M, "Notifications", "Bell shows unread count", "An outcome recorded by another user",
    "1. Observe the bell icon\n2. Open the list",
    "Unread count badge matches; notification lists the event", "Positive", "P3", "Recruiter", A_UIF, "", "notification.routes.js")
add("NTF-02", M, "Notifications", "Mark all notifications read", "Unread notifications exist",
    "1. Open the bell > Mark all read",
    "Badge clears; items shown as read after reload", "Positive", "P3", "Recruiter", A_UI, "", "notification.routes.js")

# ── MRF ──────────────────────────────────────────────────────────────────────
M = "MRF"
add("MRF-01", M, "New MRF Request", "Submit a valid New MRF Request", "Recruiter logged in",
    "1. Open /mrf\n2. Fill First Name, Last Name, Email, Role, JD Link, Budget Min 500000, Budget Max 1000000\n3. Click 'Submit Request'",
    "Toast 'MRF Request submitted successfully!'; form clears; new row in Records with MRF STATUS = PENDING; hiring manager email with the /mrf-submit link sent (inbox manual)", "Positive", "P1", "Recruiter", A_UI, EXT_MAIL, "MRF.jsx")
add("MRF-02", M, "New MRF Request", "Mandatory fields", "",
    "1. Click 'Submit Request' with an empty form",
    "'Required' under First Name, Last Name, Email, Role, JD Link, Budget Min, Budget Max; nothing submitted", "Negative", "P1", "Recruiter", A_UI, "", "MRF.jsx")
add("MRF-03", M, "New MRF Request", "Invalid hiring manager email", "",
    "1. Enter Email = 'abc@'\n2. Submit",
    "'Please enter a valid Email'", "Negative", "P2", "Recruiter", A_UI, "", "MRF.jsx")
add("MRF-04", M, "New MRF Request", "CC email ends with comma/semicolon", "",
    "1. CC Email = 'a@x.com,'\n2. Submit",
    "'CC Email should not end with comma or semicolon'", "Negative", "P2", "Recruiter", A_UI, "", "MRF.jsx validateCcEmail")
add("MRF-05", M, "New MRF Request", "Invalid CC email in list", "",
    "1. CC Email = 'a@x.com, notanemail'\n2. Submit",
    "'Please enter a valid CC email address'", "Negative", "P2", "Recruiter", A_UI, "", "MRF.jsx validateCcEmail")
add("MRF-06", M, "New MRF Request", "Budget Min below 10,000", "",
    "1. Budget Min = 5000, Budget Max = 100000\n2. Submit",
    "'Budget Min should be at least 10,000.'", "Negative", "P2", "Recruiter", A_UI, "", "MRF.jsx")
add("MRF-07", M, "New MRF Request", "Budget Max not greater than Min", "",
    "1. Budget Min = 500000, Budget Max = 500000\n2. Submit",
    "'Budget Max must be greater than Budget Min.'", "Negative", "P2", "Recruiter", A_UI, "", "MRF.jsx")
add("MRF-08", M, "New MRF Request", "Clear button resets the form", "Form partly filled",
    "1. Click 'Clear'",
    "All fields emptied; no request created", "Positive", "P3", "Recruiter", A_UI, "", "MRF.jsx")
add("MRF-09", M, "Records", "Search records", "Several MRF requests",
    "1. Type a role name in 'Search records...'",
    "Table narrows to matching rows; clearing restores all", "Positive", "P2", "Recruiter", A_UI, "", "MRF.jsx")
add("MRF-10", M, "Records", "Status filter tabs", "Requests in Pending, Manager Submitted and Closed states",
    "1. Click All / Pending / Manager Submitted / Closed",
    "Each tab shows only rows in that state", "Positive", "P2", "Recruiter", A_UI, "", "MRF.jsx")
add("MRF-11", M, "Details", "Open MRF details", "A submitted MRF",
    "1. Click a row",
    "Modal shows the request, MRF Raise Status, MRF Approval Status, the Submitted MRF Details (with AI-parsed JD summary when available) and the Approval trail; all fields read-only", "Positive", "P2", "Recruiter", A_UI, "", "MRF.jsx")
add("MRF-12", M, "Details", "Edit and save MRF details", "Details modal open",
    "1. Click Edit\n2. Change Role and Budget Max\n3. Save",
    "'MRF details updated successfully.'; MRF Raise Status select stays disabled (status is never changed by an edit)", "Positive", "P2", "Recruiter", A_UI, "", "MRF.jsx handleSaveAll")
add("MRF-13", M, "Details", "Cancel edit discards changes", "Edit mode",
    "1. Change a field\n2. Click Cancel",
    "Original values restored; nothing saved", "Positive", "P3", "Recruiter", A_UI, "", "MRF.jsx")
add("MRF-14", M, "Export", "Export MRF list CSV", "",
    "1. Click Export on Records (with a search/status filter applied)",
    "AAPNA-ATS_MRF-Requests.csv downloads containing only the filtered rows", "Positive", "P2", "Recruiter", A_UI, "", "mrf.routes.js export")
add("MRF-15", M, "Export", "Approval audit export is admin-only", "",
    "1. As Admin: check for 'Approval audit' button and download\n2. As Recruiter: check the button\n3. (API) GET /api/mrf/approval-audit/export with recruiter token",
    "Admin downloads AAPNA-ATS_MRF-Approval-Audit_90d.csv; recruiter sees no button; API returns 403", "Security", "P2", "Admin / Recruiter", A_UI, "", "mrf.routes.js")
add("MRF-16", M, "Hiring manager form", "Hiring manager submits the full MRF", "Link /mrf-submit?role=...&emailid=...&id=<request id>",
    "1. Open the link (no login)\n2. Fill all required fields; attach JD and Online Test Paper\n3. Submit",
    "'MRF Submitted Successfully!'; request shows MANAGER SUBMITTED; approval emails go to approvers", "Positive", "P1", "Hiring Manager (public)", M_EXT, "AI JD parsing; " + EXT_MAIL, "MrfSubmit.jsx")
add("MRF-17", M, "Hiring manager form", "Required fields on the public form", "",
    "1. Open /mrf-submit\n2. Submit with required fields empty",
    "'This field is required.' on each required field; email fields show 'Please enter a valid email address.'", "Negative", "P2", "Hiring Manager (public)", A_UI, "", "MrfSubmit.jsx")
add("MRF-18", M, "Hiring manager form", "Conditional 'Other' fields", "",
    "1. Set 'Requirement for the team' = Other\n2. Set 'Desired Qualification' = PG\n3. Set 'Replacement or New Role' = Replacement",
    "Matching follow-up fields appear and become required; they disappear when the trigger changes", "Positive", "P2", "Hiring Manager (public)", A_UI, "", "MrfSubmit.jsx")
add("MRF-19", M, "Hiring manager form", "JD attachment of a disallowed type", "",
    "1. Attach an .exe as the JD\n2. Submit",
    "Upload rejected with a file-type error (4xx, not a 500)", "Security", "P2", "Hiring Manager (public)", A_UI, "", "mrf.routes.js fileFilter")
add("MRF-20", M, "Approval", "Approver approves via personal link", "MRF in pending-approval; approval token for approver A (fixture mints it)",
    "1. Open /mrf/<id>/approve?token=...&action=approve\n2. Review the Requisition Summary\n3. Add a comment and Approve",
    "'Requisition Request Approved!'; approval_status = approved; decided_by = approver A; approval trail shows 'Decided'; other approvers' links are closed", "Positive", "P1", "Approver (public)", A_UIF, EXT_MAIL, "mrfApproval.service.js")
add("MRF-21", M, "Approval", "Approver rejects via personal link", "Fresh pending MRF; token for approver A",
    "1. Open the link with action=reject\n2. Add a comment and Reject",
    "'Requisition Request Declined!'; approval_status = rejected; comment stored", "Positive", "P1", "Approver (public)", A_UIF, EXT_MAIL, "MrfApprovalAction.jsx")
add("MRF-22", M, "Approval", "Second approver after a decision", "MRF decided by approver A (MRF-20)",
    "1. Approver B opens their own link",
    "Page 'Requisition already approved' with who and when (and A's comment); no second decision recorded", "Negative", "P1", "Approver (public)", A_UIF, "", "mrfApproval.service.js")
add("MRF-23", M, "Approval", "Replaced (re-issued) link", "Links re-issued after the first email (MRF-26)",
    "1. Open the OLD link",
    "'This link was replaced' / 'Please use the most recent approval email'; no decision recorded", "Negative", "P2", "Approver (public)", A_UIF, "", "MrfApprovalAction.jsx")
add("MRF-24", M, "Approval", "Invalid or tampered token", "",
    "1. Open /mrf/<id>/approve?token=garbage",
    "'Link inactive or invalid'", "Security", "P1", "Approver (public)", A_UI, "", "MrfApprovalAction.jsx")
add("MRF-25", M, "Approval", "Comment longer than 2000 characters", "Valid token",
    "1. Enter a 2001-character comment and approve",
    "Rejected: 'Comments must be at most 2000 characters.'", "Negative", "P3", "Approver (public)", A_UIF, "", "mrf.controller.js")
add("MRF-26", M, "Approval", "Re-issue approval links", "Pending MRF; Admin/Recruiter",
    "1. Details modal > re-issue approval links\n2. Repeat on an already-decided MRF",
    "Pending: new personal links sent, trail marks old ones 'Replaced by a newer link'. Decided: 'This requisition is already approved; there is nothing to re-send.'", "Positive", "P2", "Recruiter", A_UI, EXT_MAIL, "mrf.routes.js approval-links/reissue")
add("MRF-27", M, "Closure", "Close requisition needs a reason", "Approved, open MRF",
    "1. Details > 'Close requisition…'\n2. Leave Reason empty",
    "'Close it' button disabled until a reason is chosen", "Negative", "P1", "Recruiter", A_UI, "", "MRF.jsx")
add("MRF-28", M, "Closure", "Reason 'Other' needs a note", "",
    "1. Pick Reason = Other, leave the note empty",
    "'Close it' stays disabled; API also refuses: 'A note is required when the closure reason is \"other\".'", "Negative", "P2", "Recruiter", A_UI, "", "mrfClosure.service.js")
add("MRF-29", M, "Closure", "Close a requisition", "Approved, open MRF visible in the Candidate Screening JD dropdown",
    "1. Close with a valid reason\n2. Open Candidate Screening > JD dropdown",
    "'Requisition closed - it has left JD filtering.'; CLOSED tag; role no longer in the JD dropdown; approval_status unchanged", "Positive", "P1", "Recruiter", A_UI, "", "CHANGES-2026-08-26-closure-followon-mrf-lifecycle.md")
add("MRF-30", M, "Closure", "Close an already-closed requisition", "MRF closed (MRF-29)",
    "1. POST /api/mrf/<id>/close again",
    "409 'This requisition is already closed.'", "Negative", "P3", "Recruiter", A_API, "", "mrfClosure.service.js")
add("MRF-31", M, "Closure", "Re-open a requisition", "MRF closed (MRF-29)",
    "1. Details > 'Re-open requisition'\n2. Check the JD dropdown",
    "'Requisition re-opened - it is back in JD filtering.'; role back in the dropdown", "Positive", "P1", "Recruiter", A_UI, "", "MRF.jsx")
add("MRF-32", M, "Closure", "Auto-fill closes the requisition", "MRF with 1 opening; candidate at Offer",
    "1. Record offer shared > Mark accepted (OFR-02/03)\n2. Open the MRF details",
    "MRF marked filled ('all openings filled' tag); leaves JD filtering. Closing the hire as 'Did Not Join' re-opens it", "Positive", "P1", "Recruiter", A_UIF, "", "CHANGES-phase3-mrf-closure-on-offer-accepted.md")
add("MRF-33", M, "Closure", "One hire fills only one seat", "MRF with 2 openings; one candidate with accepted offer AND closed as Joined",
    "1. Check the MRF fill state",
    "MRF stays open (1 of 2 seats filled) - offer + Joined on the same journey count once", "Positive", "P2", "Recruiter", A_UIF, "", "CHANGES-2026-08-26-closure-followon-mrf-lifecycle.md")
add("MRF-34", M, "Security", "Vendor cannot export MRFs", "Vendor JWT",
    "1. GET /api/mrf/export with vendor token",
    "403 - budgets are never exposed to vendors", "Security", "P1", "Vendor", A_API, "", "mrf.routes.js")
add("MRF-35", M, "Approval", "Approval reminder emails", "Pending MRF older than the reminder threshold",
    "1. Run the reminder job (scripts/run-reminders-once.mjs)",
    "Reminder sent to approvers who have not decided; none after a decision", "Positive", "P3", "System", M_JOB, EXT_MAIL, "CHANGES-2026-09-24-mrf-approval-reminders.md")

# ── Vendor ───────────────────────────────────────────────────────────────────
M = "Vendor"
add("VND-01", M, "Vendor Dashboard", "Vendor lands on their dashboard", "Vendor account",
    "1. Log in as vendor",
    "Lands on /vendor-dashboard; KPIs 'Total Candidates', 'Added This Month', 'With Position Applied' and the Selection Rate gauge", "Positive", "P1", "Vendor", A_UI, "", "VendorDashboard.jsx")
add("VND-02", M, "Vendor Dashboard", "Vendor sees only their own candidates", "Candidates seeded for vendor A and vendor B",
    "1. Log in as vendor A\n2. Review Recent Submissions and KPIs",
    "Only vendor A's candidates appear and are counted; no vendor filter is offered", "Security", "P1", "Vendor", A_UIF, "", "utils/vendorScope.js")
add("VND-03", M, "Vendor Dashboard", "Stage column reflects the real pipeline", "Vendor candidates: one never shortlisted, one in pipeline, one closed",
    "1. Open Recent Submissions",
    "'Not in pipeline' / current stage name / 'Closed' tag respectively", "Positive", "P2", "Vendor", A_UIF, "", "CHANGES-phase3-m6-vendor.md")
add("VND-04", M, "Vendor Dashboard", "Selection rate formula", "Known counts of selected and rejected",
    "1. Hover the Selection Rate tooltip and compare with the stage breakdown",
    "Rate = Selected / (Selected + Rejected)", "Positive", "P3", "Vendor", A_UIF, "", "VendorDashboard.jsx")
add("VND-05", M, "Vendor Dashboard", "Export own submissions", "",
    "1. Click Export",
    "CSV contains only the logged-in vendor's rows", "Security", "P2", "Vendor", A_UIF, "", "VendorDashboard.jsx")
add("VND-06", M, "Vendor Dashboard", "Staff view with vendor filter", "Recruiter",
    "1. Open Vendor Dashboard as recruiter\n2. Choose a vendor in 'All Vendors'",
    "Extra 'Pending Review' KPI shown; data narrows to the selected vendor", "Positive", "P2", "Recruiter", A_UIF, "", "VendorDashboard.jsx")
add("VND-07", M, "Vendor Upload", "Upload button disabled without files", "Vendor",
    "1. Open /vendor (Upload Candidate)",
    "Dragger shows 'Supported: .pdf, .docx, .zip'; 'Upload Resumes' disabled until a file is added", "Negative", "P2", "Vendor", A_UI, "", "VendorPortal.jsx")
add("VND-08", M, "Vendor Upload", "Staff must pick a vendor", "Recruiter on /vendor",
    "1. Add a file without selecting 'On behalf of vendor'\n2. Select a vendor",
    "Tag 'Select a vendor (top right) to upload on their behalf'; button disabled. After selection: 'Uploading for: <vendor>' and button enabled", "Negative", "P2", "Recruiter", A_UI, "", "VendorPortal.jsx")
add("VND-09", M, "Vendor Upload", "Vendor uploads a resume", "Valid PDF resume",
    "1. Add file\n2. Click 'Upload Resumes'\n3. Watch Upload Status",
    "Job goes Uploaded > Processing > Saved to Database; candidate attributed to the vendor with a 90-day lock", "Positive", "P1", "Vendor", M_EXT, EXT_AI_PARSE, "VendorPortal.jsx")
add("VND-10", M, "Vendor Upload", "Upload Status filters and refresh", "Some upload jobs exist",
    "1. Filter by vendor and by status\n2. Click Refresh",
    "Counts (Total Uploads, Processing, Saved to Database, Pending Review) and table follow the filters", "Positive", "P3", "Recruiter", A_UIF, "", "VendorPortal.jsx")
add("VND-11", M, "Vendor Upload", "Second vendor cannot take a locked candidate", "Candidate locked to vendor A (within 90 days)",
    "1. Vendor B uploads the same candidate\n2. Recruiter merges the duplicate",
    "Merge keeps vendor A as owner and does not restart the 90-day window", "Security", "P1", "Vendor / Recruiter", M_EXT, EXT_AI_PARSE, "CHANGES-phase3-m6-vendor.md")
add("VND-12", M, "Vendor Upload", "Re-submission of a recently rejected candidate", "Candidate rejected < 6 months ago",
    "1. Vendor uploads the candidate again",
    "Row carries a 'Note' advisory tag about the cooling-off; resume is still queued (not blocked)", "Positive", "P3", "Vendor", M_EXT, EXT_AI_PARSE, "CHANGES-phase3-m6-vendor.md")
add("VND-13", M, "Vendor notifications", "Vendor status email per stage", "Vendor-sourced journey with an active lock",
    "1. Record an outcome at Tech 1\n2. Raise a document request\n3. Record offer shared",
    "Tech 1: vendor gets a separate status-only email (no candidate copy). Documents: vendor gets nothing. Offer: bare milestone line, no figures or dates", "Security", "P1", "Recruiter", M_EXT, EXT_MAIL, "CHANGES-phase3-m6-vendor.md")
add("VND-14", M, "Vendor isolation", "Vendor cannot widen scope via API", "Vendor A JWT",
    "1. GET /api/vendor/candidates?vendorEmail=<vendor B>",
    "Response still scoped to vendor A only", "Security", "P1", "Vendor", A_API, "", "vendorIsolation.test.js")

# ── Recruitment Analytics ────────────────────────────────────────────────────
M = "Recruitment Analytics"
add("ANL-01", M, "Overview", "Page and headline tiles load", "Seeded data",
    "1. Open /analytics",
    "Title 'Recruitment Analytics'; note 'Headline counts below are all time'; tiles Shortlisted, Rejected, On Hold, Total, Zeko Sent, Zeko Score Received", "Positive", "P1", "Recruiter", A_UI, "", "Analytics.jsx")
add("ANL-02", M, "Role Summary", "Role summary adds up", "Roles with mixed outcomes incl. a closed candidate",
    "1. Role Summary tab\n2. For each row add Shortlisted + Rejected + On Hold + Future Prospect + Closed",
    "Sum equals Total Candidates for every role (catch-all 'Closed' column)", "Positive", "P1", "Recruiter", A_UIF, "", "CHANGES-2026-08-26-candidate-closure-graceful-exit.md")
add("ANL-03", M, "Role Summary", "Export role summary", "",
    "1. Role Summary > Export",
    "AAPNA-ATS_Screening-Role-Summary.csv downloads, one row per role, incl. Closed column", "Positive", "P2", "Recruiter", A_UI, "", "Analytics.jsx")
add("ANL-04", M, "Pipeline Insights", "Hold threshold changes the tile", "Holds older and newer than 30 days",
    "1. Pipeline Insights tab\n2. Change the hold threshold 30 > 7 > 90 days",
    "'Updating analytics…' overlay; 'On hold > N days' label and count change accordingly", "Positive", "P2", "Recruiter", A_UIF, "", "Analytics.jsx")
add("ANL-05", M, "Pipeline Insights", "Stuck candidates threshold and export", "",
    "1. Change 'Stuck candidates' threshold (5/10/14/30 days)\n2. Export",
    "Heading reads 'Stuck candidates - N+ days'; table and CSV list only candidates waiting longer than N days", "Positive", "P2", "Recruiter", A_UIF, "", "Analytics.jsx")
add("ANL-06", M, "Pipeline Insights", "Rejection reasons window", "Rejections at different dates",
    "1. Change the window 7/30/90 days\n2. Export",
    "Reason counts and 'Most common stage' follow the window; empty window shows 'No rejections in this window'", "Positive", "P2", "Recruiter", A_UIF, "", "Analytics.jsx")
add("ANL-07", M, "Pipeline Insights", "Time to hire with no hires", "No closed hires",
    "1. Look at 'Median days, shortlist to hire'",
    "Shows '—' (not 0 days)", "Negative", "P3", "Recruiter", A_UIF, "", "Analytics.jsx")
add("ANL-08", M, "Pipeline Insights", "Awaiting feedback is real", "Interview held with an outstanding scorecard",
    "1. Check 'Awaiting feedback'",
    "Counts outstanding scorecards (not hard-coded 0)", "Positive", "P2", "Recruiter", A_UIF, "", "CHANGES-2026-08-26-recruitment-analytics-metric-fixes.md")
add("ANL-09", M, "Pipeline Insights", "Vendor performance and Source of hire", "Vendor-sourced and screening-sourced journeys",
    "1. Review both tables\n2. Export vendor performance",
    "Vendor rows show In pipeline / Hired / Rejected; Source rows show Submitted / In progress / Hired / Rejected / On Hold / Hire rate", "Positive", "P2", "Recruiter", A_UIF, "", "Analytics.jsx")
add("ANL-10", M, "Recruiter Insights", "Recruiter insights tab", "",
    "1. Open Recruiter Insights",
    "Per-recruiter activity loads for the same analysis window", "Positive", "P3", "Recruiter", A_UIF, "", "Analytics.jsx")
add("ANL-11", M, "Email Delivery", "Email delivery monitoring", "Emails attempted",
    "1. Open Email Delivery\n2. Click Refresh\n3. Export (role-restricted)",
    "Sent/failed emails listed; export allowed for permitted roles only", "Positive", "P3", "Recruiter", A_UIF, "", "emailTemplate.routes.js monitoring")
add("ANL-12", M, "Data accuracy", "Counts move after a pipeline action", "Baseline captured",
    "1. Reject a candidate at a stage\n2. Reload Analytics",
    "Rejected tile, role row and rejection-reason count each increase by exactly 1", "Positive", "P1", "Recruiter", A_UIF, "", "")

# ── Search Candidate ─────────────────────────────────────────────────────────
M = "Search Candidate"
add("CAND-01", M, "List", "List order and no sorting", "Seeded candidates",
    "1. Open /candidates",
    "Newest candidate (highest id) first; no sort arrows on any column", "Positive", "P2", "Recruiter", A_UIF, "", "CHANGES-2026-08-17-search-candidate-filter-simplification.md")
add("CAND-02", M, "Search", "Search by Candidate Name", "",
    "1. Enter part of a name\n2. Click Search",
    "Only matching candidates listed", "Positive", "P1", "Recruiter", A_UIF, "", "Candidates.jsx")
add("CAND-03", M, "Search", "Search by Email ID", "", "1. Enter an email\n2. Search",
    "Matching candidate listed", "Positive", "P1", "Recruiter", A_UIF, "", "Candidates.jsx")
add("CAND-04", M, "Search", "Search by Phone / Contact Number", "", "1. Enter a phone number\n2. Search",
    "Matching candidate listed", "Positive", "P1", "Recruiter", A_UIF, "", "Candidates.jsx")
add("CAND-05", M, "Search", "Two fields combine with OR", "",
    "1. Name of candidate X and email of candidate Y\n2. Search",
    "Both X and Y returned (known OR behaviour - confirm with business)", "Exploratory", "P3", "Recruiter", A_UIF, "", "CHANGES-2026-08-17-search-candidate-filter-simplification.md")
add("CAND-06", M, "Search", "Reset", "",
    "1. Observe Reset with no filter\n2. Search, then Reset",
    "Reset disabled when nothing filtered; after Reset fields clear and full list returns", "Positive", "P3", "Recruiter", A_UIF, "", "Candidates.jsx")
add("CAND-07", M, "Details", "View candidate details", "Candidate with full profile",
    "1. Click View Details",
    "Email, Phone, Location, Experience (non-zero for CVs with history), Current Company, Notice Period, Expected CTC, Education shown", "Positive", "P1", "Recruiter", A_UIF, "", "CandidateDetail.jsx")
add("CAND-08", M, "Edit", "Edit candidate", "",
    "1. Edit\n2. Change Total Experience to '5.50' and Current Location\n3. Save",
    "'Candidate details updated successfully.'; values persist after reload", "Positive", "P1", "Recruiter", A_UIF, "", "Candidates.jsx")
add("CAND-09", M, "Edit", "Invalid experience value", "",
    "1. Enter 'abc' or a negative value in experience\n2. Save",
    "Field validation error; not saved", "Negative", "P2", "Recruiter", A_UIF, "", "Candidates.jsx")
add("CAND-10", M, "Export", "Export candidates CSV", "", "1. Search\n2. Export",
    "CSV downloads with the filtered rows", "Positive", "P2", "Recruiter", A_UIF, "", "candidate.routes.js export")
add("CAND-11", M, "Referral", "Mark a candidate as referral", "",
    "1. Edit candidate\n2. Mark as referral; 'Referred by' via autocomplete\n3. Save",
    "Referral tag visible to recruiter/admin; event in referral audit log", "Positive", "P2", "Recruiter", A_UIF, "", "CHANGES-phase3-referral-candidate.md")
add("CAND-12", M, "Referral", "Only admin can remove a referral", "Referral set",
    "1. As recruiter, look for remove\n2. As admin, remove with a reason",
    "Recruiter cannot remove (API 403); admin removal needs a typed reason and is audited", "Security", "P2", "Recruiter / Admin", A_UIF, "", "candidate.routes.js")
add("CAND-13", M, "Referral", "Referral never leaks to public pages", "Referral candidate with a scorecard link",
    "1. Open the public scorecard / document page for that candidate\n2. Search the page and API responses for 'referr'",
    "No match anywhere on unauthenticated pages/responses", "Security", "P1", "Interviewer (public)", A_UIF, "", "CHANGES-phase3-referral-candidate.md")
add("CAND-14", M, "Emails", "Candidate email conversations", "Candidate with Outlook history",
    "1. Open the candidate's email conversations",
    "Thread list loads; replying sends via the recruitment mailbox", "Positive", "P3", "Recruiter", M_EXT, "MS Graph mail", "candidate.routes.js emails")
add("CAND-15", M, "Dossier", "Download complete candidate dossier", "Candidate with resume and documents",
    "1. Open the dossier download\n2. Download",
    "Pack contains the redacted profile and attachments", "Positive", "P3", "Recruiter", M_EXT, EXT_ONEDRIVE, "DossierDownloadModal.jsx")

# ── HR Manual Upload (additions) ─────────────────────────────────────────────
M = "HR Manual Upload"
add("HRU-10", M, "Upload", "Unsupported file type", "",
    "1. Upload a .txt or .exe on /hr-upload",
    "File refused (accepted: .zip, .pdf, .docx, .xlsx); no job created", "Negative", "P2", "Recruiter", A_UI, "", "HRUpload.jsx")
add("HRU-11", M, "Upload Status", "Status KPIs and refresh", "Upload jobs in several states (fixture)",
    "1. Observe Total Uploads / Processing / Saved to Database / Pending Review\n2. Click Refresh",
    "Counts match the table rows", "Positive", "P3", "Recruiter", A_UIF, "", "HRUpload.jsx")
add("HRU-12", M, "Upload Status", "Duplicate review full details", "A Duplicate_Pending_Review job (fixture)",
    "1. Click Review on the duplicate",
    "'Duplicate Review' shows Candidate, Email, File, Uploaded By; full details modal shows profile fields side by side", "Positive", "P2", "Recruiter", A_UIF, "", "HRUpload.jsx")

# ── Candidate Pipeline (additions) ───────────────────────────────────────────
M = "Candidate Pipeline"
add("PIPE-01", M, "Board", "Board columns follow configuration", "Default stages seeded",
    "1. Open /pipeline",
    "Columns in configured order; inactive 'Shortlisted' stage hidden; cards show candidate, role and status", "Positive", "P1", "Recruiter", A_UIF, "", "Pipeline.jsx")
add("PIPE-02", M, "Board", "Deep link opens the drawer", "Journey exists",
    "1. Open a card\n2. Copy URL (?candidate=...) and reload",
    "Drawer reopens on the same candidate", "Positive", "P3", "Recruiter", A_UIF, "", "PipelineDrawer.jsx")
add("PIPE-03", M, "Pause", "Pause and resume a journey", "Open journey",
    "1. Drawer > 'Pause journey'\n2. Check the card\n3. 'Resume journey'",
    "Orange 'Paused' tag on the card while paused; no closure, no email; tag removed on resume", "Positive", "P2", "Recruiter", A_UIF, "", "CHANGES-2026-08-26-closure-followon-mrf-lifecycle.md")
add("PIPE-04", M, "Closure", "Close modal requires a final status", "Open journey at any stage",
    "1. 'Close this candidate's record'\n2. Do not pick a status",
    "'Close the record' disabled; alert: 'Whether the candidate is emailed depends on the status you pick.'", "Negative", "P2", "Recruiter", A_UIF, "", "PipelineDrawer.jsx")
add("PIPE-05", M, "Closure", "Email notice per closure outcome", "",
    "1. Pick each of the 8 final statuses in turn",
    "Joined / Candidate Withdrawn / Did Not Join / Backed Out / Joined and Left: 'does not email the candidate'. Approved / Rejected / On Hold: 'emails the candidate automatically'", "Positive", "P2", "Recruiter", A_UIF, "", "PipelineDrawer.jsx")
add("PIPE-06", M, "Closure", "Withdraw a candidate from a non-Offer stage", "Journey at Technical Round 2 with a future interview booked",
    "1. Close as 'Candidate Withdrawn'",
    "Card leaves board; drawer closes; future booking cancelled (panel notified, candidate not); pipeline_status = withdrawn; candidate can be re-shortlisted (no cooling-off)", "Positive", "P1", "Recruiter", A_UIF, "", "CHANGES-2026-08-26-candidate-closure-graceful-exit.md")
add("PIPE-07", M, "Closure", "'Rejected' closure starts the 6-month cooling-off", "Journey at any stage",
    "1. Close as 'Rejected'\n2. Try to shortlist the candidate again (API)",
    "pipeline_status = rejected; re-shortlisting refused within 6 months", "Positive", "P2", "Recruiter", A_UIF, "", "CHANGES-2026-08-26-candidate-closure-graceful-exit.md")
add("PIPE-08", M, "Closure", "Closed journey suppresses actions", "Closed journey",
    "1. Open the closed journey's drawer",
    "Purple 'Closed -' tag; no Approve/Reject/Hold, Schedule, Re-invite or document actions; only 'Reopen record'", "Negative", "P2", "Recruiter", A_UIF, "", "PipelineDrawer.jsx")
add("PIPE-09", M, "Reopen", "Reopen a closed record", "Closed journey",
    "1. 'Reopen record'\n2. Try to submit without a reason\n3. Enter a reason and confirm",
    "Button disabled until a reason is typed; candidate back on the board at the stage closed on; no email; closure kept in history; MRF seats recounted", "Positive", "P1", "Recruiter", A_UIF, "", "CHANGES-2026-08-26-closure-followon-mrf-lifecycle.md")
add("PIPE-10", M, "Closure", "Booking refused on a closed journey", "Closed journey; stale tab still open",
    "1. From the stale tab, try to schedule an interview",
    "Refused with the closed-journey message; no booking created", "Security", "P3", "Recruiter", A_UIF, "", "CHANGES-2026-08-26-candidate-closure-graceful-exit.md")
add("PIPE-11", M, "Emails log", "Emails in this round log", "Outcome recorded",
    "1. Open the round's 'Emails in this round'",
    "Entry 'Outcome email > candidate · <date> · delivered / failed / not sent' matches the real send result", "Positive", "P2", "Recruiter", A_UIF, "", "PipelineDrawer.jsx")
add("PIPE-12", M, "Access", "Prototype route is module-guarded", "User without recruitment_pipeline module",
    "1. Open /candidate-pipeline-prototype",
    "Blocked like /pipeline", "Security", "P3", "Recruiter", A_UI, "", "App.jsx")

# Client Round mark-only (2026-08-25)
M = "Candidate Pipeline"
add("CLI-08", M, "Client Round (Optional)", "Client round has no scheduling", "Journey at Client Round",
    "1. Open the drawer",
    "No 'Schedule Interview' button; steps read Stage Entry > Client Meeting > Client Feedback > Approve / Reject", "Negative", "P1", "Recruiter", A_UIF, "", "CHANGES-2026-08-25-client-round-and-offer-mark-only.md")
add("CLI-09", M, "Client Round (Optional)", "Mark client meeting as held", "",
    "1. Click mark as held\n2. Enter date/time and optional client contact\n3. 'Mark as held'",
    "Step reads 'Arranged offline - held <date> with <contact>'; NO email and NO scorecard link sent to anyone", "Positive", "P1", "Recruiter", A_UIF, "", "CHANGES-2026-08-25-client-round-and-offer-mark-only.md")
add("CLI-10", M, "Client Round (Optional)", "Record client feedback", "Meeting marked held",
    "1. 'Record client feedback'\n2. Type the client's verdict\n3. 'Save feedback'",
    "Feedback shown quoted in the Client Feedback step; button not available before the meeting is marked held", "Positive", "P2", "Recruiter", A_UIF, "", "CHANGES-2026-08-25-client-round-and-offer-mark-only.md")

# ── Documents (additions) ────────────────────────────────────────────────────
M = "Documents"
add("DOC-08", M, "Candidate upload page", "Open the document upload link", "Document request raised (DOC-02); token from DB (fixture)",
    "1. Open /documents/<token> logged out",
    "'Upload your documents' with candidate/position and checklist (3 months' payslips, permanent address, government ID with DOB + father's name)", "Positive", "P1", "Candidate (public)", A_UIF, "", "DocumentUpload.jsx")
add("DOC-09", M, "Candidate upload page", "Unsupported file format", "",
    "1. Choose a .exe/.txt for an item",
    "'<file> is not an accepted format. Please upload a PDF, DOC/DOCX, JPG, PNG or ZIP.'", "Negative", "P2", "Candidate (public)", A_UIF, "", "DocumentUpload.jsx")
add("DOC-10", M, "Candidate upload page", "File over 10 MB", "",
    "1. Choose an 11 MB PDF",
    "'<file> is 11.0 MB. Please upload a file under 10 MB.'", "Negative", "P2", "Candidate (public)", A_UIF, "", "DocumentUpload.jsx")
add("DOC-11", M, "Candidate upload page", "Submit staged documents", "",
    "1. Stage files for each item\n2. Click 'Submit N documents'",
    "'Documents received - thank you.' acknowledgement; files stored under Document Collection/<name (cv-id)>/ in OneDrive", "Positive", "P1", "Candidate (public)", M_EXT, EXT_ONEDRIVE, "CHANGES-2026-08-07-candidate-pipeline-fixes.md #21")
add("DOC-12", M, "Candidate upload page", "Invalid link", "",
    "1. Open /documents/not-a-token",
    "'Link unavailable - We could not open this document request.'", "Security", "P2", "Candidate (public)", A_UI, "", "DocumentUpload.jsx")
add("DOC-13", M, "Candidate upload page", "Link closed after journey closure", "Journey closed after the request",
    "1. Open the old link",
    "'This upload link is closed'", "Security", "P2", "Candidate (public)", A_UIF, "", "CHANGES-2026-08-26-candidate-closure-graceful-exit.md")
add("DOC-14", M, "Verification", "Verify an uploaded document", "Document uploaded",
    "1. Drawer > Documents checklist > Verify",
    "Tag 'Verified'; when all items verified the segment reads 'All documents verified' and the candidate page shows 'All documents received and verified'", "Positive", "P1", "Recruiter", A_UIF, "", "PipelineDrawer.jsx")
add("DOC-15", M, "Verification", "Reject a document with a reason", "Document uploaded",
    "1. 'Reject…'\n2. Try without a reason, then enter one\n3. 'Reject & re-request'",
    "Needs a reason; tag 'Rejected - re-requested'; candidate page shows 'Please re-upload: <reason>'; re-request email to candidate only", "Positive", "P1", "Recruiter", A_UIF, EXT_MAIL, "CHANGES-phase3-m4-documents.md")
add("DOC-16", M, "Reminders", "Send reminder manually", "Request outstanding",
    "1. Click 'Send reminder'",
    "Reminder logged; counter increments only when the send succeeds", "Positive", "P3", "Recruiter", A_UIF, EXT_MAIL, "CHANGES-2026-08-07-candidate-pipeline-fixes.md #22")
add("DOC-17", M, "Reminders", "Automatic reminders", "Request outstanding > 2 days",
    "1. Run the document reminder job (daily 09:00)",
    "First reminder after 2 days, max 3; none for closed or paused journeys", "Positive", "P3", "System", M_JOB, EXT_MAIL, "jobs/documentReminder.js")

# ── Offer (additions) ────────────────────────────────────────────────────────
M = "Offer & Closure"
add("OFR-07", M, "Offer", "Joining date cannot be in the past", "Offer stage",
    "1. 'Record offer shared'\n2. Try to pick yesterday",
    "Past dates disabled in the date picker", "Negative", "P2", "Recruiter", A_UIF, "", "PipelineDrawer.jsx")
add("OFR-08", M, "Offer", "Offer progress copy", "Offer stage",
    "1. Check Awaiting Response / Accepted-Declined rows before sharing, after sharing, after decision",
    "Before: 'Awaiting the offer to be shared' / 'Not yet - the offer has not gone out'. After share: 'Awaiting the candidate's reply'. After decision: 'Candidate replied · <date>' / 'Accepted · joining <date>' or 'Declined by the candidate'", "Positive", "P3", "Recruiter", A_UIF, "", "CHANGES-2026-08-25-client-round-and-offer-mark-only.md")
add("OFR-09", M, "Closure", "Hire falls through re-opens the MRF", "MRF filled by an accepted offer",
    "1. Close the candidate as 'Did Not Join'",
    "MRF filled flag cleared; role back in JD filtering", "Positive", "P1", "Recruiter", A_UIF, "", "CHANGES-phase3-mrf-closure-on-offer-accepted.md")
add("OFR-10", M, "Closure", "Post-joining auto-close after 90 days", "Accepted offer with joining date > 90 days ago",
    "1. Run runPostJoiningAutoClose()",
    "Journey closed as 'joined'; a journey already closed as 'joined_and_left' is untouched", "Positive", "P3", "System", M_JOB, "", "jobs/offerSweep.js")

# ── Interviewer Scorecard ────────────────────────────────────────────────────
M = "Interviewer Scorecard"
add("SC-01", M, "Gate", "Scorecard link opens with context", "Tech 1 booked and marked held; scorecard token from DB (fixture)",
    "1. Open /scorecard/<token> logged out",
    "Shows CANDIDATE, POSITION, ROUND, INTERVIEWER; asks whether the interview happened ('Yes, we met - score now' / 'No - it didn't happen')", "Positive", "P1", "Interviewer (public)", A_UIF, "", "InterviewScorecard.jsx")
add("SC-02", M, "Submit", "Final rating is mandatory", "",
    "1. 'Yes, we met'\n2. Rate skills but not the final rating\n3. Submit",
    "'Please give a final rating before submitting.'", "Negative", "P1", "Interviewer (public)", A_UIF, "", "InterviewScorecard.jsx")
add("SC-03", M, "Submit", "Submit a complete scorecard", "",
    "1. Add skills with ratings and remarks, Communication, Attitude, final rating\n2. Submit (double-click)",
    "'Feedback submitted - This scorecard link is now closed.'; exactly ONE scorecard stored; recruiter sees it in 'Candidate scorecard report'", "Positive", "P1", "Interviewer (public)", A_UIF, "", "PHASE3-CODE-REVIEW-2026-08-06.md")
add("SC-04", M, "Gate", "Interview did not happen", "",
    "1. 'No - it didn't happen'",
    "'Recorded - interview did not take place'; booking marked no-show; no scorecard expected", "Positive", "P2", "Interviewer (public)", A_UIF, "", "InterviewScorecard.jsx")
add("SC-05", M, "Link state", "Used / expired / invalid link", "",
    "1. Re-open a submitted link\n2. Open an expired link\n3. Open a garbage token",
    "'Feedback submitted' (closed) / 'This scorecard link has expired' / 'Link unavailable'", "Security", "P2", "Interviewer (public)", A_UIF, "", "InterviewScorecard.jsx")
add("SC-06", M, "HR round", "HR round scorecard fields", "HR Round booking",
    "1. Open the HR round scorecard",
    "HR fields present (family background, timings, relocation, notice period, current/expected CTC, strengths, weakness, final feedback)", "Positive", "P3", "Interviewer (public)", A_UIF, "", "InterviewScorecard.jsx")

# ── Missing-data link (public) ───────────────────────────────────────────────
M = "Candidate Missing Data"
add("MJD-01", M, "Missing data", "Open the missing-data link", "Candidate saved with missing fields; token from DB (fixture)",
    "1. Open /missing-jd-upload?token=...",
    "'Complete Your Profile' with only the missing fields", "Positive", "P2", "Candidate (public)", A_UIF, "", "MissingJdUpload.jsx")
add("MJD-02", M, "Missing data", "Submit missing fields", "",
    "1. Fill the fields (attach resume .pdf/.docx if asked)\n2. Submit",
    "'Profile updated - Thank you'; candidate becomes ACTIVE in Search Candidate", "Positive", "P2", "Candidate (public)", A_UIF, "", "MissingJdUpload.jsx")
add("MJD-03", M, "Missing data", "Invalid link", "",
    "1. Open with a bad token",
    "'Link inactive or invalid'", "Security", "P3", "Candidate (public)", A_UI, "", "MissingJdUpload.jsx")
add("MJD-04", M, "Missing data", "Nothing left to fill", "Profile already complete",
    "1. Open the link again",
    "'Nothing left to fill in - Your profile is already complete.'", "Negative", "P3", "Candidate (public)", A_UIF, "", "CHANGES-empty-missing-data-token.md")

# ── Email Templates ──────────────────────────────────────────────────────────
M = "Email Templates"
add("EML-01", M, "List", "Templates list and category filter", "Templates seeded",
    "1. Open /email\n2. Filter by each category",
    "Categories: All Templates, General / Alerts, Shortlist, Interviews, Stage Outcome, Offer Letter, Rejection, Onboarding, Vendor Updates", "Positive", "P2", "Recruiter", A_UIF, "", "EmailManagement.jsx")
add("EML-02", M, "Edit", "Edit and save a template", "",
    "1. Open a template\n2. Change wording (keep placeholders)\n3. Save",
    "'Template saved successfully!'; next outcome preview uses the new wording", "Positive", "P1", "Recruiter", A_UIF, "", "EmailManagement.jsx")
add("EML-03", M, "Edit", "Mandatory placeholder removed", "",
    "1. Delete a mandatory placeholder (e.g. candidate name)\n2. Save",
    "'Cannot save: Missing mandatory placeholders.'", "Negative", "P1", "Recruiter", A_UIF, "", "EmailManagement.jsx")
add("EML-04", M, "Create", "New Template is admin-only", "",
    "1. As Admin: New Template > fill > 'Create Template'\n2. As Recruiter: look for the button",
    "Admin creates it; recruiter has no button and API POST /api/email/templates returns 403", "Security", "P2", "Admin / Recruiter", A_UIF, "", "emailTemplate.routes.js")
add("EML-05", M, "Copy", "Copy template HTML", "",
    "1. Copy HTML",
    "'Email HTML copied to clipboard.'", "Positive", "P3", "Recruiter", A_UIF, "", "EmailManagement.jsx")

# ── Settings & Configuration ─────────────────────────────────────────────────
M = "Settings & Configuration"
add("SET-01", M, "Theme", "Theme switch persists", "",
    "1. Settings > Theme: Dark\n2. Reload\n3. Switch to System and emulate OS dark/light",
    "<html data-theme> follows the choice; survives reload (localStorage ats_theme); no flash of the wrong theme", "Positive", "P3", "Any", A_UI, "", "frontend/.claude/skills/verify/SKILL.md")
add("SET-02", M, "Reminders", "Save reminder settings", "",
    "1. Change reminder values\n2. Save",
    "'Settings updated successfully!'; values persist after reload", "Positive", "P2", "Admin", A_UIF, "", "Settings.jsx")
add("SET-03", M, "Assessment", "Save assessment automation", "",
    "1. Change pass mark / auto-advance\n2. Save",
    "'Assessment automation settings updated successfully!'", "Positive", "P2", "Admin", A_UIF, "", "Settings.jsx")
add("SET-04", M, "Interviews", "Interview reminder / completion check / recording settings", "",
    "1. Change each section\n2. Save\n3. Reload",
    "Each section saves and reloads its values", "Positive", "P3", "Admin", A_UIF, "", "settings.routes.js")
add("SET-05", M, "Pipeline Configuration", "Reorder stages", "Admin",
    "1. Pipeline Configuration > 'Move stage down' on a stage\n2. Open /pipeline",
    "Board column order follows the new order", "Positive", "P2", "Admin", A_UIF, "", "PipelineConfigPanel.jsx")
add("SET-06", M, "Pipeline Configuration", "Add a rejection reason", "",
    "1. Add a reason to a stage's Reject outcome\n2. Record a Reject at that stage",
    "New reason appears in the outcome modal", "Positive", "P2", "Admin", A_UIF, "", "PipelineConfigPanel.jsx")
add("SET-07", M, "Pipeline Configuration", "Outcome Emails mapping", "",
    "1. Outcome Emails tab\n2. Map a template to stage + outcome",
    "Grid lists every stage x outcome incl. unmapped pairs; mapping used by the next outcome email preview", "Positive", "P3", "Admin", A_UIF, "", "CHANGES-phase3-m6-vendor.md")
add("SET-08", M, "Pipeline Configuration", "Email Routing (flow keys)", "",
    "1. Email Routing tab\n2. Edit a flow key's To/CC\n3. Save",
    "Saved and applied without restart; 'dynamic' / 'redirectExempt' indicators shown", "Positive", "P3", "Admin", A_UIF, "", "settings.routes.js flow-keys")
add("SET-09", M, "Pipeline Configuration", "Recruiter cannot configure the pipeline", "Recruiter",
    "1. Open Settings",
    "Pipeline Configuration panel not shown", "Security", "P2", "Recruiter", A_UIF, "", "Settings.jsx")


def build(existing):
    wb = openpyxl.Workbook()

    # ── README ──
    ws = wb.active
    ws.title = "README"
    readme = [
        ("AAPNA ATS - Master Test Cases", None),
        ("Version", "v1.0 - 2026-10-13"),
        ("Purpose", "Single catalogue of functional test cases across all ATS modules. Each case is tagged with how it will run in the Playwright suite (or why it stays manual), so the same sheet serves the demo and the formal acceptance pass."),
        ("Baseline", "Carries forward all 118 cases from 'ATS_TestCases_HRUpload_to_ZekoHR_Updated 2.xlsx' (2026-08-11 pass). Their previous result is kept in 'Previous Run (2026-08-11)'. Corrections are listed on the 'Changes vs Previous' sheet."),
        ("Target environment", "Local dev: frontend (Vite :5173) + backend (:5000) + dedicated test database (to be agreed). External services are NOT called."),
        ("", ""),
        ("Automation legend", ""),
        (A_UI, "Playwright drives the browser; needs only a running app and seeded users."),
        (A_UIF, "Playwright UI test that needs a fixture first - e.g. a journey created via API, or a token (scorecard / documents / MRF approval) read from the test DB."),
        (A_API, "Playwright API request test (security / authorisation checks with no UI)."),
        (M_EXT, "Depends on an external service (AI parser/search, Zeko, MS Graph mail/Teams/OneDrive, Turnstile). Skipped in the local suite; run manually on staging."),
        (M_JOB, "Depends on a cron/background job; run manually by invoking the job."),
        (M_EXP, "Open product question - needs a decision before it can have a pass/fail."),
        ("", ""),
        ("Email checks", "Locally, outcome/notification emails cannot be delivered. Automated cases assert the state change and the 'Emails in this round' log; real inbox delivery is a manual staging check (column 'External Dependency' = MS Graph mail)."),
        ("Priority", "P1 = core flow / security, must pass for the demo. P2 = important validation. P3 = secondary."),
        ("Status values", "Not Run, Pass, Fail, Blocked, Skipped"),
    ]
    for i, (k, v) in enumerate(readme, start=1):
        ws.cell(row=i, column=1, value=k)
        if v is not None:
            ws.cell(row=i, column=2, value=v)
    ws["A1"].font = Font(bold=True, size=14)
    for r in range(2, len(readme) + 1):
        ws.cell(row=r, column=1).font = Font(bold=True)
        ws.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=r, column=1).alignment = Alignment(vertical="top")
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 110

    # ── Test Cases ──
    tc = wb.create_sheet("Test Cases")
    tc.append(HEADERS)
    rows = existing + [r[:13] + ["New", "Not Run", "", ""] for r in N]
    for r in rows:
        tc.append(r)

    head_fill = PatternFill("solid", fgColor="1F4E78")
    thin = Side(style="thin", color="BFBFBF")
    for c in tc[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head_fill
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    widths = [10, 20, 24, 34, 30, 44, 50, 12, 8, 16, 22, 26, 30, 16, 10, 24, 40]
    for i, w in enumerate(widths, start=1):
        tc.column_dimensions[get_column_letter(i)].width = w
    for row in tc.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
    tc.freeze_panes = "B2"
    tc.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{tc.max_row}"
    tc.row_dimensions[1].height = 32

    last = tc.max_row
    status_col = get_column_letter(HEADERS.index("Status") + 1)
    dv = DataValidation(type="list", formula1='"Not Run,Pass,Fail,Blocked,Skipped"', allow_blank=True)
    tc.add_data_validation(dv)
    dv.add(f"{status_col}2:{status_col}{last}")
    rng = f"A2:{get_column_letter(len(HEADERS))}{last}"
    for val, color in (("Pass", "E2EFDA"), ("Fail", "FCE4D6"), ("Blocked", "FFF2CC")):
        tc.conditional_formatting.add(rng, FormulaRule(formula=[f'${status_col}2="{val}"'], fill=PatternFill("solid", fgColor=color)))

    # ── Summary ──
    sm = wb.create_sheet("Summary")
    modules = []
    for r in rows:
        if r[1] not in modules:
            modules.append(r[1])
    autos = [A_UI, A_UIF, A_API, M_EXT, M_JOB, M_EXP]
    hdr = ["Module", "Total", "P1"] + autos + ["Automatable", "Automatable %", "Pass", "Fail", "Not Run"]
    sm.append(hdr)
    B = "'Test Cases'!$B$2:$B$%d" % last
    I = "'Test Cases'!$I$2:$I$%d" % last
    K = "'Test Cases'!$K$2:$K$%d" % last
    O = "'Test Cases'!$O$2:$O$%d" % last
    for i, m in enumerate(modules, start=2):
        sm.cell(row=i, column=1, value=m)
        sm.cell(row=i, column=2, value=f'=COUNTIF({B},$A{i})')
        sm.cell(row=i, column=3, value=f'=COUNTIFS({B},$A{i},{I},"P1")')
        for j, a in enumerate(autos):
            sm.cell(row=i, column=4 + j, value=f'=COUNTIFS({B},$A{i},{K},"{a}")')
        sm.cell(row=i, column=10, value=f"=D{i}+E{i}+F{i}")
        sm.cell(row=i, column=11, value=f"=IF(B{i}=0,0,J{i}/B{i})")
        sm.cell(row=i, column=12, value=f'=COUNTIFS({B},$A{i},{O},"Pass")')
        sm.cell(row=i, column=13, value=f'=COUNTIFS({B},$A{i},{O},"Fail")')
        sm.cell(row=i, column=14, value=f'=COUNTIFS({B},$A{i},{O},"Not Run")')
    tot = len(modules) + 2
    sm.cell(row=tot, column=1, value="Total").font = Font(bold=True)
    for col in range(2, 15):
        L = get_column_letter(col)
        if col == 11:
            sm.cell(row=tot, column=col, value=f"=IF(B{tot}=0,0,J{tot}/B{tot})")
        else:
            sm.cell(row=tot, column=col, value=f"=SUM({L}2:{L}{tot - 1})")
        sm.cell(row=tot, column=col).font = Font(bold=True)
    for r in range(2, tot + 1):
        sm.cell(row=r, column=11).number_format = "0%"
    for c in sm[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head_fill
        c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
    sm.column_dimensions["A"].width = 28
    for col in range(2, 15):
        sm.column_dimensions[get_column_letter(col)].width = 14
    sm.row_dimensions[1].height = 45
    sm.freeze_panes = "B2"

    # ── Changes vs Previous ──
    ch = wb.create_sheet("Changes vs Previous")
    ch.append(["TC ID", "What changed", "Why"])
    why_map = {
        "OFR": "Offer stage became record-only on 2026-08-25 (CHANGES-2026-08-25-client-round-and-offer-mark-only.md). Approve/Reject/Hold no longer exist on Offer.",
        "T1-12": "Open 'Verify' item resolved from code.", "T1-13": "Open 'Verify' item resolved from code.",
        "T3-08": "Skip control now exists in the drawer UI.", "CLI-07": "Skip control now exists in the drawer UI.",
    }
    for tcid, o in OVERRIDES.items():
        changed = [k for k in ("scenario", "steps", "expected", "type", "automation", "pre") if k in o]
        if not changed and not o.get("remarks_add"):
            continue
        what = ("Updated: " + ", ".join(changed)) if changed else "Remark added"
        why = why_map.get(tcid) or why_map.get(tcid.split("-")[0]) or (o.get("remarks_add") or "")
        ch.append([tcid, what, why])
    ch.append([])
    ch.append(["Feedback (2026-08-11)", "Interviewer name in scheduled email", "Fixed - CHANGES-2026-08-07 #20 (re-test T1-07/T1-09)"])
    ch.append(["Feedback (2026-08-11)", "Submit button on document upload", "Already present; acknowledgement added - #21 (now DOC-11)"])
    ch.append(["Feedback (2026-08-11)", "Automatic document reminders", "Already present (daily 09:00, after 2 days, max 3) - #22 (now DOC-17)"])
    ch.append([])
    ch.append(["New modules", ", ".join(sorted({r[1] for r in N})), "Not covered by the previous sheet"])
    for c in ch[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head_fill
    ch.column_dimensions["A"].width = 22
    ch.column_dimensions["B"].width = 50
    ch.column_dimensions["C"].width = 110
    for row in ch.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")

    wb.active = 0
    wb.save(OUT)
    return rows


if __name__ == "__main__":
    existing = existing_case_rows(SRC)
    ids = [r[0] for r in existing] + [r[0] for r in N]
    dupes = {i for i in ids if ids.count(i) > 1}
    assert not dupes, f"duplicate ids: {dupes}"
    for r in N:
        assert len(r) == 13, (r[0], len(r))
        assert r[10] in (A_UI, A_UIF, A_API, M_EXT, M_JOB, M_EXP), r[0]
    rows = build(existing)
    from collections import Counter
    print("existing:", len(existing), "new:", len(N), "total:", len(rows))
    print(Counter(r[10] for r in rows))
    print(Counter(r[1] for r in rows))
